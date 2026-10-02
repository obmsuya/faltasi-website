from openpyxl import load_workbook

from django.db import transaction

from .models import (
    Department,
    RequestCategory,
    RequestField,
)


REQUIRED_HEADERS = [
    "Organization",
    "Department",
    "Category",
    "Description",
    "Keywords",
    "Field Name",
    "Question",
    "Field Type",
    "Required",
    "Order",
]


@transaction.atomic
def import_categories_from_excel(file_path, organization):
    """
    Import request categories and customer questions.

    The organization is selected by the administrator.
    The Organization column in Excel is informational only.

    Existing categories and questions are updated rather than duplicated.

    The complete operation runs inside one database transaction.
    If an error occurs, the import is rolled back.
    """

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    try:
        if "Category Questions" not in workbook.sheetnames:
            raise ValueError(
                "The Excel file must contain a sheet named "
                "'Category Questions'."
            )

        worksheet = workbook["Category Questions"]

        # -----------------------------------------------------
        # Read headers
        # -----------------------------------------------------

        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            None,
        )

        if not first_row:
            raise ValueError(
                "The Excel file is empty."
            )

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in first_row
        ]

        missing_headers = [
            header
            for header in REQUIRED_HEADERS
            if header not in headers
        ]

        if missing_headers:
            raise ValueError(
                "Missing required Excel columns: "
                + ", ".join(missing_headers)
            )

        column_index = {
            header: index
            for index, header in enumerate(headers)
        }

        # -----------------------------------------------------
        # Statistics
        # -----------------------------------------------------

        categories_created = 0
        categories_updated = 0
        categories_unchanged = 0

        fields_created = 0
        fields_updated = 0
        fields_unchanged = 0

        rows_processed = 0
        rows_skipped = 0

        created_categories = []
        processed_fields = []

        counted_categories = set()

        # -----------------------------------------------------
        # Process rows
        # -----------------------------------------------------

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):

            # Ignore completely empty rows
            if not any(row):
                rows_skipped += 1
                continue

            rows_processed += 1

            def get_value(column_name):
                index = column_index[column_name]

                if index >= len(row):
                    return None

                return row[index]

            department_name = get_value("Department")
            category_name = get_value("Category")
            description = get_value("Description")
            keywords = get_value("Keywords")
            field_name = get_value("Field Name")
            question = get_value("Question")
            field_type = get_value("Field Type")
            required = get_value("Required")
            order = get_value("Order")

            # -------------------------------------------------
            # Validate required values
            # -------------------------------------------------

            if not department_name:
                raise ValueError(
                    f"Row {row_number}: Department is required."
                )

            if not category_name:
                raise ValueError(
                    f"Row {row_number}: Category is required."
                )

            if not field_name:
                raise ValueError(
                    f"Row {row_number}: Field Name is required."
                )

            if not question:
                raise ValueError(
                    f"Row {row_number}: Question is required."
                )

            # -------------------------------------------------
            # Clean values
            # -------------------------------------------------

            department_name = str(
                department_name
            ).strip()

            category_name = str(
                category_name
            ).strip()

            field_name = str(
                field_name
            ).strip()

            question = str(
                question
            ).strip()

            description = (
                str(description).strip()
                if description
                else ""
            )

            keywords = (
                str(keywords).strip()
                if keywords
                else ""
            )

            field_type = (
                str(field_type).strip()
                if field_type
                else "text"
            )

            # -------------------------------------------------
            # Find department
            # -------------------------------------------------

            department = Department.objects.filter(
                organization=organization,
                name=department_name,
                is_active=True,
            ).first()

            if not department:
                raise ValueError(
                    f"Row {row_number}: Department "
                    f"'{department_name}' does not exist for "
                    f"'{organization.name}'."
                )

            # -------------------------------------------------
            # Find or create category
            # -------------------------------------------------

            category, category_created = (
                RequestCategory.objects.get_or_create(
                    organization=organization,
                    name=category_name,
                    defaults={
                        "department": department,
                        "description": description,
                        "keywords": keywords,
                    },
                )
            )

            category_changed = False

            if category.department_id != department.id:
                category.department = department
                category_changed = True

            if category.description != description:
                category.description = description
                category_changed = True

            if category.keywords != keywords:
                category.keywords = keywords
                category_changed = True

            if category_changed:
                category.save()

            # -------------------------------------------------
            # Count each category once
            # -------------------------------------------------

            if category.id not in counted_categories:

                counted_categories.add(category.id)

                if category_created:
                    categories_created += 1
                    created_categories.append(category)

                elif category_changed:
                    categories_updated += 1

                else:
                    categories_unchanged += 1

            # -------------------------------------------------
            # Convert required value
            # -------------------------------------------------

            is_required = (
                str(required).strip().lower()
                in ["yes", "true", "1"]
            )

            # -------------------------------------------------
            # Convert order
            # -------------------------------------------------

            try:
                field_order = int(order)
            except (TypeError, ValueError):
                field_order = 1

            # -------------------------------------------------
            # Find or create question
            # -------------------------------------------------

            field, field_created = (
                RequestField.objects.get_or_create(
                    category=category,
                    name=field_name,
                    defaults={
                        "question": question,
                        "field_type": field_type,
                        "is_required": is_required,
                        "order": field_order,
                        "is_active": True,
                    },
                )
            )

            # -------------------------------------------------
            # Detect question changes
            # -------------------------------------------------

            field_changed = False

            if field.question != question:
                field.question = question
                field_changed = True

            if field.field_type != field_type:
                field.field_type = field_type
                field_changed = True

            if field.is_required != is_required:
                field.is_required = is_required
                field_changed = True

            if field.order != field_order:
                field.order = field_order
                field_changed = True

            if not field.is_active:
                field.is_active = True
                field_changed = True

            if field_changed:
                field.save()

            processed_fields.append(field)

            if field_created:
                fields_created += 1

            elif field_changed:
                fields_updated += 1

            else:
                fields_unchanged += 1

        return {
            "categories": created_categories,
            "fields": processed_fields,

            "categories_created": categories_created,
            "categories_updated": categories_updated,
            "categories_unchanged": categories_unchanged,

            "fields_created": fields_created,
            "fields_updated": fields_updated,
            "fields_unchanged": fields_unchanged,

            "rows_processed": rows_processed,
            "rows_skipped": rows_skipped,
        }

    finally:
        workbook.close()