
import csv
import os
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_EXTRACT_SIZE = 50 * 1024 * 1024

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".csv",
    ".txt",
    ".md",
}


class DocumentExtractionError(Exception):
    """Raised when a document cannot be safely extracted."""


def extract_document_text(file_path):
    """
    Extract text from a supported document.

    Returns:
        dict: {
            "text": extracted text,
            "page_count": page or worksheet count where applicable,
            "file_extension": extension
        }
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            f"Unsupported file type: {extension or 'no extension'}"
        )

    try:
        if not path.is_file():
            raise DocumentExtractionError("The document file was not found.")

        if path.stat().st_size > MAX_EXTRACT_SIZE:
            raise DocumentExtractionError(
                "The document exceeds the 50 MB processing limit."
            )

        if path.stat().st_size == 0:
            raise DocumentExtractionError("The document is empty.")

        if extension == ".pdf":
            text, page_count = _extract_pdf(path)

        elif extension == ".docx":
            text, page_count = _extract_docx(path), None

        elif extension == ".xlsx":
            text, page_count = _extract_xlsx(path)

        elif extension == ".csv":
            text, page_count = _extract_csv(path), None

        else:
            text, page_count = _extract_text_file(path), None

        text = text.strip()

        if not text:
            raise DocumentExtractionError(
                "No readable text was found. The file may be scanned "
                "or contain only images. OCR may be required."
            )

        return {
            "text": text,
            "page_count": page_count,
            "file_extension": extension.lstrip("."),
        }

    except DocumentExtractionError:
        raise

    except Exception as exc:
        # Avoid returning internal paths or sensitive parser details.
        raise DocumentExtractionError(
            "The document could not be read. Check that it is a valid, "
            "unprotected file of the declared type."
        ) from exc


def _extract_pdf(path):
    reader = PdfReader(str(path), strict=True)

    if reader.is_encrypted:
        raise DocumentExtractionError(
            "Password-protected PDFs are not supported."
        )

    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")

    return "\n\n".join(pages), len(reader.pages)


def _extract_docx(path):
    document = WordDocument(str(path))
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)


def _extract_xlsx(path):
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )

    parts = []

    try:
        for worksheet in workbook.worksheets:
            parts.append(f"Worksheet: {worksheet.title}")

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    parts.append(" | ".join(values))

        return "\n".join(parts), len(workbook.worksheets)

    finally:
        workbook.close()


def _extract_csv(path):
    # utf-8-sig supports CSV files with a UTF-8 byte-order mark.
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = csv.reader(file)
        return "\n".join(
            " | ".join(cell.strip() for cell in row)
            for row in rows
        )


def _extract_text_file(path):
    # Strict decoding avoids silently corrupting product names or prices.
    return path.read_text(encoding="utf-8-sig")
